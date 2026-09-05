import os, sqlite3, uuid
from pathlib import Path
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_from_directory
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from google import genai

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env", override=True)
UPLOAD_DIR = BASE_DIR / "uploads"; UPLOAD_DIR.mkdir(exist_ok=True)
DB_PATH = BASE_DIR / "creatorchat.db"
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
TEXT_MODEL = os.getenv("GEMINI_TEXT_MODEL", "gemini-3.6-flash").strip()
client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY", "creatorchat-local-secret")
ALLOWED = {"pdf","docx","doc","xlsx","xls","pptx","ppt","txt","csv","md"}

def get_db():
    c = sqlite3.connect(DB_PATH); c.row_factory = sqlite3.Row; return c

def init_db():
    c = get_db(); c.executescript("""
    CREATE TABLE IF NOT EXISTS chats(id INTEGER PRIMARY KEY AUTOINCREMENT,title TEXT,user_message TEXT,ai_message TEXT,created_at TEXT);
    CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY AUTOINCREMENT,filename TEXT,stored_name TEXT,extracted_text TEXT,created_at TEXT);
    CREATE TABLE IF NOT EXISTS schedules(id INTEGER PRIMARY KEY AUTOINCREMENT,title TEXT,platform TEXT,schedule_time TEXT,notes TEXT,created_at TEXT);
    CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT,description TEXT,created_at TEXT);
    """); c.commit(); c.close()
init_db()

def ai_text(prompt, system="You are CreatorChat House, a helpful AI assistant for YouTubers and digital creators. Support English, Tamil and Tanglish. Give practical, ready-to-use answers."):
    if not client: raise RuntimeError("GEMINI_API_KEY is not configured.")
    r = client.models.generate_content(model=TEXT_MODEL, contents=f"{system}\n\n{prompt}")
    return getattr(r, "text", "") or "Gemini returned an empty response."

def extract_text(path):
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if ext == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(str(path)).paragraphs)
        if ext in {".xlsx",".xls"}:
            from openpyxl import load_workbook
            wb=load_workbook(str(path),read_only=True,data_only=True); out=[]
            for ws in wb.worksheets:
                out.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True): out.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(out)
        if ext == ".pptx":
            from pptx import Presentation
            out=[]
            for i,s in enumerate(Presentation(str(path)).slides,1):
                out.append(f"[Slide {i}]")
                for sh in s.shapes:
                    if hasattr(sh,"text"): out.append(sh.text)
            return "\n".join(out)
        if ext in {".txt",".csv",".md"}: return path.read_text(encoding="utf-8",errors="ignore")
        if ext in {".doc",".ppt"}: return "Legacy file uploaded. Convert to DOCX/PPTX for automatic text extraction."
    except Exception as e: return f"Extraction error: {e}"
    return ""

@app.route("/")
def index(): return render_template("index.html")

@app.get("/api/health")
def health():
    return jsonify(ok=True, provider="Google Gemini", api_configured=bool(GEMINI_API_KEY), text_model=TEXT_MODEL, image_generation=False, image_prompt_generation=True)

@app.post("/api/chat")
def chat():
    d=request.get_json(silent=True) or {}; msg=(d.get("message") or "").strip()
    if not msg: return jsonify(error="Message is required"),400
    try:
        ans=ai_text(msg); c=get_db(); c.execute("INSERT INTO chats(title,user_message,ai_message,created_at) VALUES(?,?,?,?)",(msg[:80],msg,ans,datetime.now().isoformat(timespec="seconds"))); c.commit(); c.close(); return jsonify(message=ans)
    except Exception as e: app.logger.exception("Gemini chat error"); return jsonify(error=f"Gemini request failed: {e}"),500

IMAGE_PROMPT_SYSTEM = """You are an expert AI image-prompt engineer for professional creators. Convert the user's simple description into ONE detailed, production-ready prompt for an external AI image generator. Understand Tamil, Tanglish and English. Preserve the user's core idea, but improve visual specificity. Include subject, appearance, action/pose, environment, composition, camera angle, lighting, mood, color palette, depth, visual hierarchy, lens/camera feel, requested style, aspect ratio and useful negative constraints. Make it cinematic and easy to paste into an image generator. Return ONLY the final prompt, with no explanation and no markdown heading."""

@app.post("/api/image-prompt")
def image_prompt():
    d=request.get_json(silent=True) or {}; desc=(d.get("description") or "").strip(); style=(d.get("style") or "").strip(); ratio=(d.get("ratio") or "16:9").strip()
    if not desc: return jsonify(error="Describe the image you want"),400
    request_text=f"""Turn this simple image idea into a polished AI image prompt.\n\nUSER DESCRIPTION:\n{desc}\n\nPREFERRED STYLE: {style or 'cinematic and professional'}\nASPECT RATIO: {ratio}\n\nCreate one complete prompt suitable for an external image generator."""
    try: return jsonify(prompt=ai_text(request_text, IMAGE_PROMPT_SYSTEM))
    except Exception as e: app.logger.exception("image prompt"); return jsonify(error=f"Image prompt generation failed: {e}"),500

TOOL_PROMPTS={"hooks":"Generate 10 highly clickable YouTube hooks for this topic. Mix curiosity, surprise, challenge, story and result-driven hooks.","titles":"Generate 15 natural, clickable and non-misleading YouTube titles for this topic.","script":"Write a complete YouTube script with Hook, Intro, Main Points, examples, transitions, CTA and Outro.","seo":"Create an SEO package: title, description, keywords/tags, hashtags and pinned-comment idea.","ideas":"Generate 20 practical content ideas around this topic, with a one-line angle for each.","thumbnail":"Create 8 thumbnail concepts. For each give visual scene, emotion, composition, text suggestion and an image prompt."}
@app.post("/api/tool")
def tool():
    d=request.get_json(silent=True) or {}; name=d.get("tool","hooks"); topic=(d.get("input") or "").strip()
    if name not in TOOL_PROMPTS or not topic: return jsonify(error="Tool and input are required"),400
    try: return jsonify(message=ai_text(TOOL_PROMPTS[name]+"\n\nTopic:\n"+topic))
    except Exception as e: return jsonify(error=f"Gemini request failed: {e}"),500

@app.post("/api/upload")
def upload():
    f=request.files.get("file")
    if not f or not f.filename: return jsonify(error="No file selected"),400
    if "." not in f.filename or f.filename.rsplit(".",1)[1].lower() not in ALLOWED: return jsonify(error="Unsupported file type"),400
    original=secure_filename(f.filename); stored=f"{uuid.uuid4().hex}_{original}"; path=UPLOAD_DIR/stored; f.save(path); text=extract_text(path)
    c=get_db(); cur=c.execute("INSERT INTO documents(filename,stored_name,extracted_text,created_at) VALUES(?,?,?,?)",(original,stored,text,datetime.now().isoformat(timespec="seconds"))); c.commit(); doc_id=cur.lastrowid; c.close()
    return jsonify(id=doc_id,filename=original,extracted_chars=len(text),message="File uploaded and processed.")

@app.get("/api/documents")
def documents():
    c=get_db(); rows=c.execute("SELECT id,filename,created_at FROM documents ORDER BY id DESC").fetchall(); c.close(); return jsonify(documents=[dict(r) for r in rows])

@app.post("/api/document-chat")
def document_chat():
    d=request.get_json(silent=True) or {}; doc_id=d.get("document_id"); q=(d.get("question") or "").strip(); c=get_db(); row=c.execute("SELECT filename,extracted_text FROM documents WHERE id=?",(doc_id,)).fetchone(); c.close()
    if not row: return jsonify(error="Document not found"),404
    if not q: return jsonify(error="Question is required"),400
    prompt=f"Answer using this uploaded document. If the answer is not present, say so clearly.\nDocument: {row['filename']}\nCONTENT:\n{(row['extracted_text'] or '')[:120000]}\nQUESTION:\n{q}"
    try: return jsonify(message=ai_text(prompt))
    except Exception as e: return jsonify(error=f"Gemini request failed: {e}"),500

@app.get("/api/chats")
def chats():
    c=get_db(); rows=c.execute("SELECT * FROM chats ORDER BY id DESC LIMIT 100").fetchall(); c.close(); return jsonify(chats=[dict(r) for r in rows])

@app.route("/api/schedules",methods=["GET","POST"])
def schedules():
    c=get_db()
    if request.method=="GET": rows=c.execute("SELECT * FROM schedules ORDER BY id DESC").fetchall(); c.close(); return jsonify(schedules=[dict(r) for r in rows])
    d=request.get_json(silent=True) or {}; title=(d.get("title") or "").strip(); t=(d.get("schedule_time") or "").strip()
    if not title or not t: c.close(); return jsonify(error="Title and schedule time are required"),400
    cur=c.execute("INSERT INTO schedules(title,platform,schedule_time,notes,created_at) VALUES(?,?,?,?,?)",(title,d.get("platform",""),t,d.get("notes",""),datetime.now().isoformat(timespec="seconds"))); c.commit(); row=c.execute("SELECT * FROM schedules WHERE id=?",(cur.lastrowid,)).fetchone(); c.close(); return jsonify(dict(row))

@app.route("/api/projects",methods=["GET","POST"])
def projects():
    c=get_db()
    if request.method=="GET": rows=c.execute("SELECT * FROM projects ORDER BY id DESC").fetchall(); c.close(); return jsonify(projects=[dict(r) for r in rows])
    d=request.get_json(silent=True) or {}; name=(d.get("name") or "").strip()
    if not name: c.close(); return jsonify(error="Project name is required"),400
    cur=c.execute("INSERT INTO projects(name,description,created_at) VALUES(?,?,?)",(name,d.get("description",""),datetime.now().isoformat(timespec="seconds"))); c.commit(); row=c.execute("SELECT * FROM projects WHERE id=?",(cur.lastrowid,)).fetchone(); c.close(); return jsonify(dict(row))

@app.get("/api/search")
def search():
    q=(request.args.get("q") or "").strip(); c=get_db(); rows=c.execute("SELECT * FROM chats WHERE user_message LIKE ? OR ai_message LIKE ? ORDER BY id DESC LIMIT 50",(f"%{q}%",f"%{q}%")).fetchall(); c.close(); return jsonify(results=[dict(r) for r in rows])

@app.get("/uploads/<path:name>")
def files(name): return send_from_directory(UPLOAD_DIR,name)

if __name__=="__main__": app.run(debug=True)
